"""Unit tests for document parsers (Task 3.1, R9.2).

Verifies text and structure extraction across:
- Plain text (.txt)
- Markdown (.md)
- HTML (.html)
- PDF (.pdf)
- DOCX (.docx)
- Parser registry and error handling
"""

from __future__ import annotations

import pytest

from packages.domain.knowledge import ElementType
from packages.knowledge.parsers.base import HeadingStack
from packages.knowledge.parsers.text import PlainTextParser


class TestHeadingStack:
    """Test heading hierarchy stack operations."""

    def test_push_and_pop_hierarchy(self) -> None:
        stack = HeadingStack()
        h_path, sec = stack.push(1, "Guide")
        assert h_path == ("Guide",)
        assert sec == "Guide"

        h_path, sec = stack.push(2, "Auth")
        assert h_path == ("Guide", "Auth")
        assert sec == "Auth"

        h_path, sec = stack.push(3, "OAuth2")
        assert h_path == ("Guide", "Auth", "OAuth2")
        assert sec == "OAuth2"

        # Sibling heading at level 2 pops prior level 2 and level 3
        h_path, sec = stack.push(2, "Billing")
        assert h_path == ("Guide", "Billing")
        assert sec == "Billing"

        # Level 1 pops everything
        h_path, sec = stack.push(1, "API Reference")
        assert h_path == ("API Reference",)
        assert sec == "API Reference"


class TestPlainTextParser:
    """Test plain text parser structure extraction."""

    def test_parse_headings_and_paragraphs(self) -> None:
        content = """Title Document
===============

This is an introductory paragraph.

Section One
-----------
Paragraph in section one.

### Subsection 1.1
Paragraph in subsection.
"""
        parser = PlainTextParser()
        doc = parser.parse(content, filename="manual.txt")

        assert doc.title == "Title Document"
        assert len(doc.elements) >= 5

        # Headings verification
        headings = doc.get_elements_by_type(ElementType.HEADING)
        assert len(headings) == 3
        assert headings[0].content == "Title Document"
        assert headings[0].level == 1
        assert headings[0].heading_path == ("Title Document",)

        assert headings[1].content == "Section One"
        assert headings[1].level == 2
        assert headings[1].heading_path == ("Title Document", "Section One")

        assert headings[2].content == "Subsection 1.1"
        assert headings[2].level == 3
        assert headings[2].heading_path == ("Title Document", "Section One", "Subsection 1.1")

        # Paragraphs verification with heading_path
        paragraphs = doc.get_elements_by_type(ElementType.PARAGRAPH)
        assert len(paragraphs) == 3
        assert paragraphs[0].heading_path == ("Title Document",)
        assert paragraphs[1].heading_path == ("Title Document", "Section One")
        assert paragraphs[2].heading_path == ("Title Document", "Section One", "Subsection 1.1")

    def test_parse_lists(self) -> None:
        content = """# Setup Checklist

- Step 1: Install dependencies
- Step 2: Configure environment
  * Nested item

1. First ordered task
2. Second ordered task
"""
        parser = PlainTextParser()
        doc = parser.parse(content)

        lists = doc.get_elements_by_type(ElementType.LIST_ITEM)
        assert len(lists) == 5
        assert lists[0].content == "Step 1: Install dependencies"
        assert lists[0].metadata["list_type"] == "bullet"
        assert lists[0].heading_path == ("Setup Checklist",)

        assert lists[3].content == "First ordered task"
        assert lists[3].metadata["list_type"] == "ordered"

    def test_parse_empty_content(self) -> None:
        parser = PlainTextParser()
        doc = parser.parse("", filename="empty.txt")
        assert doc.title == "Empty"
        assert len(doc.elements) == 0
        assert doc.text == ""


class TestMarkdownParser:
    """Test Markdown parser structure extraction."""

    def test_parse_frontmatter_and_headings(self) -> None:
        from packages.knowledge.parsers.markdown import MarkdownParser

        content = """---
title: System Architecture Guide
author: Antigravity Team
---

# Overview
This document describes the enterprise architecture.

## Ingestion Pipeline
Ingestion runs asynchronously.

### Parsers
Parsers convert raw formats.
"""
        parser = MarkdownParser()
        doc = parser.parse(content, filename="arch.md")

        assert doc.title == "System Architecture Guide"
        headings = doc.get_elements_by_type(ElementType.HEADING)
        assert len(headings) == 3
        assert headings[0].content == "Overview"
        assert headings[0].level == 1
        assert headings[0].heading_path == ("Overview",)

        assert headings[1].content == "Ingestion Pipeline"
        assert headings[1].level == 2
        assert headings[1].heading_path == ("Overview", "Ingestion Pipeline")

        assert headings[2].content == "Parsers"
        assert headings[2].level == 3
        assert headings[2].heading_path == ("Overview", "Ingestion Pipeline", "Parsers")

    def test_parse_tables_and_code_blocks(self) -> None:
        from packages.knowledge.parsers.markdown import MarkdownParser

        content = """# Data Specifications

| Parameter | Type | Default |
|---|---|---|
| timeout | int | 30 |
| retries | int | 3 |

```python
def configure():
    return True
```

> Important notice about timeouts.
"""
        parser = MarkdownParser()
        doc = parser.parse(content)

        tables = doc.get_elements_by_type(ElementType.TABLE)
        assert len(tables) == 1
        table_meta = tables[0].metadata.get("table_data", {})
        assert table_meta["headers"] == ["Parameter", "Type", "Default"]
        assert len(table_meta["rows"]) == 2
        assert table_meta["rows"][0] == ["timeout", "int", "30"]

        codes = doc.get_elements_by_type(ElementType.CODE_BLOCK)
        assert len(codes) == 1
        assert "def configure():" in codes[0].content
        assert codes[0].metadata.get("language") == "python"

        quotes = doc.get_elements_by_type(ElementType.BLOCKQUOTE)
        assert len(quotes) == 1
        assert "Important notice" in quotes[0].content

    def test_parse_nested_lists(self) -> None:
        from packages.knowledge.parsers.markdown import MarkdownParser

        content = """# Roadmap

- Phase 1: Ingestion
  - Subtask 1.1
- Phase 2: Retrieval
"""
        parser = MarkdownParser()
        doc = parser.parse(content)

        lists = doc.get_elements_by_type(ElementType.LIST_ITEM)
        assert len(lists) == 3
        assert lists[0].content == "Phase 1: Ingestion"
        assert lists[1].content == "Subtask 1.1"
        assert lists[1].metadata["list_depth"] == 2


class TestHTMLParser:
    """Test HTML parser structure extraction."""

    def test_parse_semantic_tags_and_noise_stripping(self) -> None:
        from packages.knowledge.parsers.html import HTMLParser

        html = """<!DOCTYPE html>
<html>
<head>
    <title>Customer Knowledge Base</title>
    <script>console.log('malicious/noisy script');</script>
    <style>.hide { display: none; }</style>
</head>
<body>
    <nav><a href="/home">Home</a></nav>
    <h1>Knowledge Article</h1>
    <p>Welcome to customer support.</p>
    <h2>Account Reset</h2>
    <p>Steps to reset password:</p>
    <ul>
        <li>Click forgot password</li>
        <li>Enter email address</li>
    </ul>
    <table>
        <tr><th>Role</th><th>Permission</th></tr>
        <tr><td>Admin</td><td>Full Access</td></tr>
    </table>
    <pre><code>curl -X GET /status</code></pre>
    <footer>Copyright 2026</footer>
</body>
</html>
"""
        parser = HTMLParser()
        doc = parser.parse(html, filename="article.html")

        assert doc.title == "Customer Knowledge Base"
        # Scripts and nav must be stripped
        assert "console.log" not in doc.text
        assert "Home" not in doc.text
        assert "Copyright" not in doc.text

        headings = doc.get_elements_by_type(ElementType.HEADING)
        assert len(headings) == 2
        assert headings[0].content == "Knowledge Article"
        assert headings[0].level == 1
        assert headings[1].content == "Account Reset"
        assert headings[1].level == 2
        assert headings[1].heading_path == ("Knowledge Article", "Account Reset")

        lists = doc.get_elements_by_type(ElementType.LIST_ITEM)
        assert len(lists) == 2
        assert lists[0].content == "Click forgot password"
        assert lists[0].heading_path == ("Knowledge Article", "Account Reset")

        tables = doc.get_elements_by_type(ElementType.TABLE)
        assert len(tables) == 1
        assert tables[0].metadata["table_data"]["headers"] == ["Role", "Permission"]
        assert tables[0].metadata["table_data"]["rows"][0] == ["Admin", "Full Access"]

        code_blocks = doc.get_elements_by_type(ElementType.CODE_BLOCK)
        assert len(code_blocks) == 1
        assert "curl -X GET" in code_blocks[0].content


class TestPDFParser:
    """Test PDF parser with in-memory synthetic PDF."""

    def test_parse_pdf_pages_and_metadata(self) -> None:
        import io

        from pypdf import PdfWriter

        from packages.knowledge.parsers.pdf import PDFParser

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)

        # Set title and metadata in PDF
        writer.add_metadata({
            "/Title": "Synthetic PDF Specification",
            "/Author": "Antigravity Engineering",
        })

        buf = io.BytesIO()
        writer.write(buf)
        buf.seek(0)

        parser = PDFParser()
        doc = parser.parse(buf.getvalue(), filename="spec.pdf")

        assert doc.title == "Synthetic PDF Specification"
        assert doc.metadata["page_count"] == 2
        assert doc.metadata["author"] == "Antigravity Engineering"
        assert doc.mime_type == "application/pdf"

    def test_parse_corrupt_pdf_raises_document_parsing_error(self) -> None:
        from packages.knowledge.parsers.base import DocumentParsingError
        from packages.knowledge.parsers.pdf import PDFParser

        parser = PDFParser()
        with pytest.raises(DocumentParsingError) as exc_info:
            parser.parse(b"NOT_A_VALID_PDF_HEADER", filename="broken.pdf")
        assert "Failed to read PDF document" in str(exc_info.value)


class TestDOCXParser:
    """Test DOCX parser with in-memory synthetic DOCX."""

    def test_parse_docx_structure_and_tables(self) -> None:
        import io

        import docx

        from packages.knowledge.parsers.docx import DOCXParser

        d = docx.Document()
        d.core_properties.title = "Enterprise Operations Guide"
        d.core_properties.author = "Reliability Team"

        d.add_heading("System Operations", level=1)
        d.add_paragraph("Introductory instructions.")

        d.add_heading("Emergency Procedures", level=2)
        d.add_paragraph("First step of recovery.", style="List Bullet")
        d.add_paragraph("Second step of recovery.", style="List Bullet")

        tbl = d.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = "Severity"
        tbl.rows[0].cells[1].text = "SLA"
        tbl.rows[1].cells[0].text = "P1"
        tbl.rows[1].cells[1].text = "15 min"

        buf = io.BytesIO()
        d.save(buf)
        buf.seek(0)

        parser = DOCXParser()
        doc = parser.parse(buf.getvalue(), filename="ops.docx")

        assert doc.title == "Enterprise Operations Guide"
        assert doc.metadata["author"] == "Reliability Team"

        headings = doc.get_elements_by_type(ElementType.HEADING)
        assert len(headings) == 2
        assert headings[0].content == "System Operations"
        assert headings[0].level == 1
        assert headings[1].content == "Emergency Procedures"
        assert headings[1].level == 2
        assert headings[1].heading_path == ("System Operations", "Emergency Procedures")

        lists = doc.get_elements_by_type(ElementType.LIST_ITEM)
        assert len(lists) == 2
        assert lists[0].content == "First step of recovery."
        assert lists[0].heading_path == ("System Operations", "Emergency Procedures")

        tables = doc.get_elements_by_type(ElementType.TABLE)
        assert len(tables) == 1
        assert tables[0].metadata["table_data"]["headers"] == ["Severity", "SLA"]
        assert tables[0].metadata["table_data"]["rows"][0] == ["P1", "15 min"]

    def test_parse_corrupt_docx_raises_document_parsing_error(self) -> None:
        from packages.knowledge.parsers.base import DocumentParsingError
        from packages.knowledge.parsers.docx import DOCXParser

        parser = DOCXParser()
        with pytest.raises(DocumentParsingError) as exc_info:
            parser.parse(b"NOT_A_VALID_ZIP_OR_DOCX", filename="corrupt.docx")
        assert "Failed to read DOCX document" in str(exc_info.value)


class TestParserRegistry:
    """Test parser registry resolution and error handling."""

    def test_resolve_by_mime_type(self) -> None:
        from packages.knowledge.parsers.docx import DOCXParser
        from packages.knowledge.parsers.html import HTMLParser
        from packages.knowledge.parsers.markdown import MarkdownParser
        from packages.knowledge.parsers.pdf import PDFParser
        from packages.knowledge.parsers.registry import get_parser
        from packages.knowledge.parsers.text import PlainTextParser

        assert isinstance(get_parser(mime_type="text/plain"), PlainTextParser)
        assert isinstance(get_parser(mime_type="text/markdown"), MarkdownParser)
        assert isinstance(get_parser(mime_type="text/html"), HTMLParser)
        assert isinstance(get_parser(mime_type="application/pdf"), PDFParser)
        assert isinstance(
            get_parser(
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            DOCXParser,
        )

    def test_resolve_by_filename_extension(self) -> None:
        from packages.knowledge.parsers.docx import DOCXParser
        from packages.knowledge.parsers.html import HTMLParser
        from packages.knowledge.parsers.markdown import MarkdownParser
        from packages.knowledge.parsers.pdf import PDFParser
        from packages.knowledge.parsers.registry import get_parser
        from packages.knowledge.parsers.text import PlainTextParser

        assert isinstance(get_parser(filename="notes.txt"), PlainTextParser)
        assert isinstance(get_parser(filename="README.md"), MarkdownParser)
        assert isinstance(get_parser(filename="index.html"), HTMLParser)
        assert isinstance(get_parser(filename="report.pdf"), PDFParser)
        assert isinstance(get_parser(filename="proposal.docx"), DOCXParser)

    def test_unsupported_document_type_raises(self) -> None:
        from packages.knowledge.parsers.base import UnsupportedDocumentTypeError
        from packages.knowledge.parsers.registry import get_parser, parse_document

        with pytest.raises(UnsupportedDocumentTypeError):
            get_parser(mime_type="audio/mpeg", filename="audio.mp3")

        with pytest.raises(UnsupportedDocumentTypeError):
            parse_document(b"binary", filename="archive.tar.gz")

    def test_convenience_parse_document(self) -> None:
        from packages.knowledge.parsers.registry import parse_document

        doc = parse_document("# Direct Title\n\nDirect body paragraph.", filename="quick.md")
        assert doc.title == "Direct Title"
        assert len(doc.elements) == 2
        assert doc.elements[0].element_type == ElementType.HEADING
        assert doc.elements[1].element_type == ElementType.PARAGRAPH


class TestParsedDocumentHelpers:
    """Test ParsedDocument structure helpers."""

    def test_get_sections_and_heading_hierarchy(self) -> None:
        from packages.domain.knowledge import DocumentElement, ParsedDocument

        elements = [
            DocumentElement(
                element_type=ElementType.HEADING,
                content="Chapter 1",
                level=1,
                heading_path=("Chapter 1",),
                section="Chapter 1",
            ),
            DocumentElement(
                element_type=ElementType.PARAGRAPH,
                content="Paragraph in chapter 1.",
                heading_path=("Chapter 1",),
                section="Chapter 1",
            ),
            DocumentElement(
                element_type=ElementType.HEADING,
                content="Section 1.1",
                level=2,
                heading_path=("Chapter 1", "Section 1.1"),
                section="Section 1.1",
            ),
            DocumentElement(
                element_type=ElementType.PARAGRAPH,
                content="Paragraph in section 1.1.",
                heading_path=("Chapter 1", "Section 1.1"),
                section="Section 1.1",
            ),
        ]
        doc = ParsedDocument(title="Book", elements=elements)

        hierarchy = doc.get_heading_hierarchy()
        assert hierarchy == [(1, "Chapter 1"), (2, "Section 1.1")]

        sections = doc.get_sections()
        assert len(sections) == 2
        assert sections[0].title == "Chapter 1"
        assert len(sections[0].elements) == 2  # heading + paragraph
        assert sections[1].title == "Section 1.1"
        assert len(sections[1].elements) == 2

        assert "Chapter 1" in doc.to_text()
        assert "Paragraph in section 1.1." in doc.to_text()
